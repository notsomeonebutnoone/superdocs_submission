from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


SOURCE = Path(__file__).with_name("atlas-field-systems-proposal.md")
OUTPUT = Path(__file__).with_name("atlas-field-systems-detailed-proposal.docx")


def add_markdown_line(document: Document, line: str) -> None:
    if line.startswith("# "):
        paragraph = document.add_heading(line[2:], level=0)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif line.startswith("## "):
        document.add_heading(line[3:], level=1)
    elif line.startswith("### "):
        document.add_heading(line[4:], level=2)
    elif line.startswith("- "):
        document.add_paragraph(line[2:], style="List Bullet")
    elif line:
        paragraph = document.add_paragraph()
        pieces = line.split("**")
        for index, piece in enumerate(pieces):
            paragraph.add_run(piece).bold = index % 2 == 1


document = Document()
section = document.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)

styles = document.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Title"].font.name = "Aptos Display"
styles["Title"].font.size = Pt(24)
styles["Heading 1"].font.name = "Aptos Display"
styles["Heading 1"].font.size = Pt(16)

for source_line in SOURCE.read_text(encoding="utf-8").splitlines():
    add_markdown_line(document, source_line.strip())

footer = section.footer.paragraphs[0]
footer.text = "Northstar Workflow Studio | Synthetic evaluation document"
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

document.save(OUTPUT)
print(OUTPUT)
